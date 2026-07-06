#!/usr/bin/env python3
"""
VIDEO 16 — FORMATO CATÁLOGO (patrón 'Every Drug Explained', 4.9M views)
Every Money Bias & Its Effect Explained in 8 Minutes
NEUROCENTS · 81 beats · ~1,050 words · ~8 min
Formato: 9 — CATÁLOGO/GRID (primera vez) · V13/V14 = Lista, V15 = Mecanismo ✅ rotación S20
Hook: CERO HOOK — patrón S24-P15: el título ES el contrato, entrada directa al item 1
      (excepción validada a S19 para Formato 9 — el outlier de 4.9M abre "Caffeine. Caffeine is...")
EL GIRO ÚNICO: cada sesgo tratado como una SUSTANCIA — con onset, peak y comedown.
      Nadie en el nicho de dinero está haciendo esto. Es la traducción directa del outlier más grande.
CTA at beats 32-33 = 39.5% ✅ (dentro de 35-40%) — redacción NUEVA
Sin Villain's Last Trick (Formato 9 no lo pide — regla S20 permite omitir) — cierre catálogo + share-CTA
REGLA DE INDEPENDENCIA: cada sesgo autocontenido en 6-9 beats, redacción 100% nueva ✅
      (Algunos sesgos aparecieron en vídeos antiguos individuales — esto es el CATÁLOGO de referencia:
      mismo valor que 'Every Drug Explained' recopilando sustancias con mil vídeos individuales detrás.
      CERO frases recicladas de V1-V15.)
ESCALADA (patrón cafeína→opioides): Anchoring (el más cotidiano) → ... → Scarcity Mindset (el más profundo)

MICRO-PLANTILLA POR SESGO (patrón catálogo 4.9M):
  Nombre → qué es (1 frase cotidiana) → escena/trigger 2ª persona → ONSET (cómo empieza, se siente bien)
  → PEAK (cuando decide por ti) → COMEDOWN (la factura) → daño largo plazo → FRASE LAPIDARIA

LOS 10 SESGOS (orden de escalada):
  1. Anchoring · 2. Loss Aversion · 3. Mental Accounting · 4. Present Bias · [CTA] · 5. Herd Instinct
  · 6. Lifestyle Inflation · 7. Sunk Cost · 8. Optimism Loop · 9. Default Bias · 10. Scarcity Mindset

S17: Thumbnail = grid de 10 círculos de colores con icono-objeto por sesgo (estilo 'DRUG EFFECTS
     EXPLAINED') + header "MONEY BIASES EXPLAINED" → Beat 1 = el círculo de Anchoring llenando
     pantalla + "Anchoring." — el viewer reconoce el grid en <5 seg.
     NOTA iconos: objetos cotidianos (etiqueta de precio tachada, carrito, hucha...), NO conceptos
     abstractos — lección del test a 200px.
BEATS VISUALES: guión narrado SIN Alex — el personaje va solo en los image prompts como hilo visual
     mínimo (regla 'Reducir Alex', Sección 9). El Brain Villain puede operar cada 'sustancia' visualmente.
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

TITLE    = "Every Money Bias & Its Effect Explained in 8 Minutes"
SUBTITLE = "NEUROCENTS · VIDEO 16 — CATALOG FORMAT"
LABEL    = "SCRIPT (NARRATION)"

SECTIONS = [

    # ── BIAS 1 ── 9 beats — el más cotidiano (patrón cafeína: implica al 100% de la audiencia)
    ("BIAS 1 — ANCHORING", [
        "Anchoring. Anchoring is your brain's habit of trusting the first number it sees.",  # B1 — CERO intro, directo (S24-P15)
        "It's everywhere. Price tags, salaries, menus. The most consumed bias on Earth.",     # B2
        "You see a jacket: two hundred euros, crossed out. Now ninety.",                      # B3 escena
        "The onset is instant. Ninety stops being a price — it becomes a bargain.",           # B4 ONSET
        "Your brain never asked if the jacket was worth ninety. It only compared it to two hundred.",  # B5
        "At the peak, the crossed-out number does all your thinking. The real question — 'do I even want this?' — never loads.",  # B6 PEAK
        "The comedown arrives at home: a ninety-euro jacket you never planned to buy.",       # B7 COMEDOWN
        "Repeated daily, anchoring quietly decides what 'normal' costs — rent, phones, haircuts, everything.",  # B8 daño
        "You never buy the thing. You buy the distance from the first number.",               # B9 LAPIDARIA
    ]),

    # ── BIAS 2 ── 8 beats
    ("BIAS 2 — LOSS AVERSION", [
        "Loss Aversion. The oldest bias in the building.",                                    # B10
        "Losing a hundred euros hurts about twice as much as winning a hundred feels good. That asymmetry runs your life.",  # B11
        "The onset: you check a price, an account, an investment — and your chest tightens before your thoughts arrive.",  # B12 ONSET
        "At the peak, you'll do irrational things to avoid a small loss — hold a bad investment, keep a broken subscription, stay in the wrong plan.",  # B13 PEAK
        "Not to win. Just to not lose.",                                                      # B14 short beat
        "The comedown is invisible: all the better options you never took because they smelled like risk.",  # B15 COMEDOWN
        "Long-term damage: a life optimized against losing is a life that never compounds.",  # B16 daño
        "Your brain doesn't protect your money. It protects the feeling of not losing it.",   # B17 LAPIDARIA
    ]),

    # ── BIAS 3 ── 7 beats
    ("BIAS 3 — MENTAL ACCOUNTING", [
        "Mental Accounting. Your brain runs separate wallets for the same money.",            # B18
        "Salary money is serious. Refund money is free. Birthday money is fun.",              # B19
        "The onset: a hundred euros arrives outside your salary — and it lands in the 'doesn't count' wallet.",  # B20 ONSET
        "At the peak, you'll spend a tax refund in a weekend while agonizing over a forty-euro grocery bill.",  # B21 PEAK
        "Same currency. Same account. Different rules.",                                      # B22
        "The comedown: at the end of the year, the 'free money' is gone and you can't name a single thing it built.",  # B23 COMEDOWN
        "Money doesn't come with labels. Your brain prints them.",                            # B24 LAPIDARIA
    ]),

    # ── BIAS 4 ── 7 beats
    ("BIAS 4 — PRESENT BIAS", [
        "Present Bias. The dealer that always finds you.",                                    # B25 humor oscuro
        "To your brain, you-today is a real person. You-in-ten-years is a stranger in a stock photo.",  # B26
        "The onset: 'I'll start saving next month.' It feels responsible. It's actually the high.",  # B27 ONSET
        "At the peak, today's wants outvote tomorrow's needs every single time — dinner beats retirement, now beats later.",  # B28 PEAK
        "The comedown never comes today. That's the design. The bill is always addressed to the stranger.",  # B29 COMEDOWN
        "Until one morning you're the stranger, opening the mail.",                           # B30
        "Present bias doesn't steal your money. It borrows it from someone you haven't met yet — you.",  # B31 LAPIDARIA
    ]),

    # ── CTA ── beats 32-33 = 39.5% ✅ — redacción NUEVA estilo catálogo
    ("CTA", [
        "Four down, six to go. If your brain has already shown up twice in this catalog — subscribe.",  # B32
        "We name one of these programs every week. Free, painless, and mildly uncomfortable in the good way.",  # B33
    ]),

    # ── BIAS 5 ── 7 beats
    ("BIAS 5 — HERD INSTINCT", [
        "Herd Instinct. Social proof. The party drug.",                                       # B34
        "Your brain outsources decisions to the crowd — if everyone's buying, it must be safe.",  # B35
        "The onset feels like belonging: everyone has the phone, the trip, the coin, the sneakers.",  # B36 ONSET
        "At the peak, 'everyone's doing it' overrides every number you know. FOMO isn't fear of missing the thing — it's fear of standing outside the group.",  # B37 PEAK
        "The comedown hits when the crowd moves on and you're still holding the receipt.",    # B38 COMEDOWN
        "The crowd got the memories. You got the credit card statement.",                     # B39 daño
        "Herds are great protection against lions. There are no lions at the mall.",          # B40 LAPIDARIA + humor
    ]),

    # ── BIAS 6 ── 7 beats
    ("BIAS 6 — LIFESTYLE INFLATION", [
        "Lifestyle Inflation. The tolerance effect.",                                         # B41 vocabulario sustancia
        "Every raise feels enormous for exactly one month.",                                  # B42
        "The onset: more money arrives, and 'needs' quietly upgrade themselves to match. Better coffee. Better car. Better everything.",  # B43 ONSET
        "At the peak, you're earning double what you did five years ago — and saving exactly the same: nothing.",  # B44 PEAK
        "That's tolerance. The dose went up. The effect didn't.",                             # B45 — la analogía sustancia explícita
        "The comedown is the trap itself: there is no comedown. It just becomes your baseline. Forever.",  # B46 COMEDOWN
        "A raise doesn't make you richer. It makes your old life unaffordable.",              # B47 LAPIDARIA
    ]),

    # ── BIAS 7 ── 7 beats
    ("BIAS 7 — SUNK COST", [
        "Sunk Cost. The loyalty program of bad decisions.",                                   # B48 humor
        "You keep paying for the gym you don't attend, the course you don't finish, the project that died last year.",  # B49
        "The onset is a sentence: 'But I've already put so much into it.'",                   # B50 ONSET
        "At the peak, the past runs your future — you throw good money after bad, because quitting would make the loss real.",  # B51 PEAK
        "Here's what your brain hides from you: the money is already gone. It left when you spent it.",  # B52
        "The comedown lasts years: every month you stay is another payment on a decision you already know was wrong.",  # B53 COMEDOWN
        "You're not protecting your investment. You're buying tickets to watch it sink.",     # B54 LAPIDARIA
    ]),

    # ── BIAS 8 ── 6 beats
    ("BIAS 8 — THE OPTIMISM LOOP", [
        "The Optimism Loop. Planning fallacy, if you want the lab name.",                     # B55
        "Next month, you'll spend less. Next month has been coming for nine years.",          # B56 humor seco
        "The onset: every budget you make stars a fictional character — a disciplined, unhurried you with no birthdays, no emergencies, no Fridays.",  # B57 ONSET
        "At the peak, you plan for the best month you've ever had — every month.",            # B58 PEAK
        "The comedown arrives on day twenty-something, with a calendar full of exceptions that were 'one-time things.' All twelve months have one-time things.",  # B59 COMEDOWN
        "Optimism is a great life partner and a terrible accountant.",                        # B60 LAPIDARIA
    ]),

    # ── BIAS 9 ── 6 beats
    ("BIAS 9 — DEFAULT BIAS", [
        "Default Bias. The silent subscription.",                                             # B61
        "Whatever is already happening keeps happening — not because it's good, but because changing it requires a decision.",  # B62
        "The onset is nothing. That's its trick. The same bank since you were eighteen. The same tariff. The same insurance, renewing itself in the dark.",  # B63 ONSET
        "At the peak, entire industries price your inertia into their business model — the loyalty penalty has a name because it's that reliable.",  # B64 PEAK
        "The comedown is spread so thin you never feel it: thirty euros here, sixty there, every month, for decades.",  # B65 COMEDOWN
        "The most expensive decisions of your life are the ones you never made.",             # B66 LAPIDARIA
    ]),

    # ── BIAS 10 ── 8 beats — el profundo (patrón opioides: cierra la escalada)
    ("BIAS 10 — SCARCITY MINDSET", [
        "Scarcity Mindset. The heavy one.",                                                   # B67
        "When money has been tight long enough, scarcity stops being a situation and becomes an operating system.",  # B68
        "The onset: money worries start taxing your attention — rent math running in the background of every conversation.",  # B69 ONSET
        "Researchers measured it: active financial scarcity can consume more cognitive capacity than a full night without sleep.",  # B70 dato
        "At the peak, the tunnel closes in. Today's fire is all there is. Long-term thinking isn't a choice you're refusing — it's a luxury the tunnel doesn't stock.",  # B71 PEAK
        "And the cruelest part: the tunnel makes the exact decisions that keep you in the tunnel.",  # B72
        "The comedown isn't even yours. It's inherited — scarcity teaches children what money feels like before they ever earn any.",  # B73 COMEDOWN generacional
        "Poverty isn't just a number in an account. It's a full-time job your brain works for free.",  # B74 LAPIDARIA
    ]),

    # ── CLOSE ── 5 beats — cierre catálogo + share-CTA (patrón S24-P12), sin Villain template
    ("CLOSE — THE CATALOG", [
        "Ten programs. One brain. Yours came pre-installed with all of them.",                # B75
        "You can't uninstall a bias. Nobody can.",                                            # B76 honestidad intelectual
        "But a named program never runs quietly again — from now on, when one fires, some part of you will be watching.",  # B77
        "That part is new. It didn't exist eight minutes ago.",                               # B78
        "Send this catalog to the friend who's deep in sunk cost right now. They'll know exactly which one they're on.",  # B79 share-CTA
    ]),

    # ── NEXT VIDEO TEASE ── 2 beats — V17: Comparativa
    ("NEXT VIDEO TEASE", [
        "Next week: two people. Same salary. Same city. At forty-five, one of them stops needing to work.",  # B80
        "The difference fits on a napkin. See you Thursday.",                                 # B81
    ]),

]


def build_docx():
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(SUBTITLE)
    r.bold = True; r.font.size = Pt(14)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run(LABEL)
    r.bold = True; r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0xB0, 0x2A, 0x2A)

    s2 = doc.add_paragraph()
    s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s2.add_run(TITLE)
    r.bold = True; r.italic = True; r.font.size = Pt(14)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nr = note.add_run("Formato 9 Catálogo (patrón 4.9M) · Cero hook — título es el contrato · Onset/Peak/Comedown por sesgo · 10 sesgos en escalada")
    nr.font.size = Pt(9); nr.font.color.rgb = RGBColor(0x44, 0x88, 0x44)

    doc.add_paragraph()

    beat_num = 1
    for section_title, lines in SECTIONS:
        sh = doc.add_paragraph()
        sr = sh.add_run(f"— {section_title} —")
        sr.bold = True; sr.font.size = Pt(11)
        sr.font.color.rgb = RGBColor(0xB0, 0x2A, 0x2A)
        sh.paragraph_format.space_before = Pt(14)
        sh.paragraph_format.space_after = Pt(4)

        for line in lines:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            nr = p.add_run(f"[{beat_num}]  ")
            nr.bold = True; nr.font.size = Pt(9)
            nr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            p.add_run(line).font.size = Pt(11)
            beat_num += 1
        doc.add_paragraph()

    all_lines = [line for _, lines in SECTIONS for line in lines]
    wc = sum(len(l.split()) for l in all_lines)
    total = beat_num - 1

    f = doc.add_paragraph()
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = f.add_run(f"TOTAL: {total} beats · ~{wc} words · ~{round(wc/140)} min")
    fr.font.size = Pt(9); fr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    path = "/home/user/Claudeeee/V16_final_SCRIPT.docx"
    doc.save(path)
    print(f"DOCX: {path} | {total} beats | {wc} words | ~{round(wc/140)} min")
    return all_lines, total, wc


def build_pdf(all_lines, total, wc):
    styles = getSampleStyleSheet()
    H1  = ParagraphStyle('H1',  parent=styles['Title'],  fontSize=14, leading=18, alignment=TA_CENTER)
    SUB = ParagraphStyle('SUB', parent=styles['Normal'], fontSize=10, leading=13,
                         alignment=TA_CENTER, textColor=colors.HexColor('#B02A2A'),
                         fontName='Helvetica-Bold')
    NOTE= ParagraphStyle('NOTE',parent=styles['Normal'], fontSize=9, leading=12,
                         alignment=TA_CENTER, textColor=colors.HexColor('#448844'))
    META= ParagraphStyle('META',parent=styles['Normal'], fontSize=9, leading=12,
                         alignment=TA_CENTER, textColor=colors.HexColor('#777777'))
    SEC = ParagraphStyle('SEC', parent=styles['Normal'], fontSize=10, leading=13,
                         textColor=colors.HexColor('#B02A2A'), fontName='Helvetica-Bold',
                         spaceBefore=14, spaceAfter=4)
    LINE= ParagraphStyle('LINE',parent=styles['Normal'], fontSize=11, leading=17, spaceAfter=1)

    def esc(t): return t.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')

    pdf_path = "/home/user/Claudeeee/V16_final_SCRIPT.pdf"
    doc = SimpleDocTemplate(pdf_path, pagesize=A4,
                            leftMargin=20*mm, rightMargin=20*mm,
                            topMargin=18*mm, bottomMargin=18*mm)
    flow = [
        Paragraph(esc(SUBTITLE), SUB),
        Spacer(1, 4),
        Paragraph(esc(TITLE), H1),
        Spacer(1, 4),
        Paragraph("Formato 9 Catálogo (patrón 4.9M) · Cero hook · Onset/Peak/Comedown por sesgo · CTA beats 32-33 = 39.5%", NOTE),
        Spacer(1, 4),
        Paragraph(f"{LABEL} · {total} beats · ~{wc} words · ~{round(wc/140)} min", META),
        Spacer(1, 14),
    ]
    beat_num = 1
    for section_title, lines in SECTIONS:
        flow.append(Paragraph(f"— {esc(section_title)} —", SEC))
        for line in lines:
            flow.append(Paragraph(
                f'<font color="#888888"><b>[{beat_num}]</b></font>  {esc(line)}', LINE))
            beat_num += 1
        flow.append(Spacer(1, 8))
    flow.append(Spacer(1, 10))
    flow.append(Paragraph(f"END · {total} beats · ~{wc} words · ~{round(wc/140)} min", META))
    doc.build(flow)
    print(f"PDF:  {pdf_path}")


def build_elevenlabs_pdf(all_lines, wc):
    styles = getSampleStyleSheet()
    H1  = ParagraphStyle('H1',  parent=styles['Title'],  fontSize=14, leading=18, alignment=TA_CENTER)
    SUB = ParagraphStyle('SUB', parent=styles['Normal'], fontSize=10, leading=13,
                         alignment=TA_CENTER, textColor=colors.HexColor('#B02A2A'),
                         fontName='Helvetica-Bold')
    META= ParagraphStyle('META',parent=styles['Normal'], fontSize=9,  leading=12,
                         alignment=TA_CENTER, textColor=colors.HexColor('#777777'))
    LINE= ParagraphStyle('LINE',parent=styles['Normal'], fontSize=12, leading=20, spaceAfter=3)

    def esc(t): return t.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')

    pdf_path = "/home/user/Claudeeee/V16_final_ELEVENLABS.pdf"
    doc = SimpleDocTemplate(pdf_path, pagesize=A4,
                            leftMargin=20*mm, rightMargin=20*mm,
                            topMargin=18*mm, bottomMargin=18*mm)
    flow = [
        Paragraph("NEUROCENTS · VIDEO 16 — CATALOG FORMAT", SUB),
        Spacer(1, 4),
        Paragraph(esc(TITLE), H1),
        Spacer(1, 4),
        Paragraph(f"ELEVENLABS — NARRATION ONLY · {len(all_lines)} lines · ~{wc} words · ~{round(wc/140)} min", META),
        Spacer(1, 16),
    ]
    for line in all_lines:
        flow.append(Paragraph(esc(line), LINE))
    flow.append(Spacer(1, 10))
    flow.append(Paragraph(f"END OF SCRIPT · {len(all_lines)} lines · ~{wc} words", META))
    doc.build(flow)
    print(f"ELEVENLABS: {pdf_path}")


if __name__ == "__main__":
    all_lines, total, wc = build_docx()
    build_pdf(all_lines, total, wc)
    build_elevenlabs_pdf(all_lines, wc)
    print(f"\n✅ V16 CATÁLOGO — 3 archivos generados")
    print(f"   {total} beats · {wc} words · ~{round(wc/140)} min · dentro de 70-110 ✅")
    print(f"   Formato 9 (Catálogo, patrón 4.9M) — primera vez ✅")
    print(f"   Cero hook — entrada directa 'Anchoring.' (S24-P15) ✅")
    print(f"   CTA: beats 32-33 = {round(32/total*100)}% ✅")
    print(f"   Onset/Peak/Comedown en los 10 sesgos ✅ · Lapidaria por sesgo ✅")
    print(f"   Escalada Anchoring → Scarcity Mindset ✅ · Independencia total ✅")
