#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the Etiqueta Real / Aceite de Oliva production spreadsheet.
Real 3-5s beats (~9-11 words each), full original script, 400+ beats."""
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from full_script import SECTIONS
from new_shots_urls import NEW_URLS

# ---------- SHOT LIBRARY (deduplicated stock clips used across the video) ----------
SHOTS = {
    "buyer_choosing":   ("The buyer chooses olive oil in the supermarket", 151138),
    "receipt_check":    ("Person Checking Bill Or Receipt For Supermarket Shopping Basket Full Of Basic Food Items", 687809),
    "shelf_oils":       ("Close up of olive oil and other cooking oil bottles on a shelf in a supermarket", 3472735),
    "premium_bottle":   ("An Exquisite Representation of the Finest Olive Oil: A Harvest Reserve Bottle", 6264187),
    "magnifier_doc":    ("Close-Up of a Magnifying Glass on an Aged Document Highlighting Detailed Text and Graphs", 8534240),
    "lab_pipette":      ("Scientist using pipette 4k", 1864061),
    "aisle_dolly":      ("A smooth out-of-focus camera movement travels down a brightly lit retail grocery store aisle", 7917819),
    "bottles_row":      ("Bottles of Cooking Oil", 5177810),
    "network_anim":     ("Central avatar appearing and expanding, sending connecting lines and nodes, showing company network", 7618572),
    "yellow_cap":       ("Close-up of a Bottle with Yellow Cap", 3811676),
    "euros_rotating":   ("Rotating Around Euros", 6285),
    "three_bottles":    ("Three glass bottles with handles containing olive oil, closed with cork stopper, with a bowl of olives", 5465997),
    "opening_cap":      ("Hand opening a bottle cap with a bottle opener", 5474555),
    "bottling_line":    ("Automated Beverage Bottling Line", 5540725),
    "conveyor_bottles":  ("Bottled Drink Production Line", 5540686),
    "olive_grove_aerial": ("Aerial View of Olive Grove Landscape", 6956087),
    "olive_grove_sunset": ("Vast Olive Grove at Sunset", 6163120),
    "warehouse_barrels": ("Red and Blue Industrial Barrels Stored in an Outdoor Warehouse Yard", 7249724),
    "two_bottles_compare": ("Two bottles of olive oil and extra virgin olive", 2310359),
    "port_containers":  ("Aerial View of a Busy Container Port", 5883923),
    "world_map_port":   ("Appearing world map overlay animating above port terminal, glowing arcs routing cargo toward cranes", 7558377),
    "generic_shelf":    ("Fixed clip of full grocery store shelves, filled with brightly coloured products", 6181813),
    "hand_reaching_jar": ("A hand reaches for and inspects a jar of olives in a grocery store aisle", 6977454),
    "andalusia_street": ("Traditional whitewashed houses with balconies and barred windows line a narrow cobblestone street, Andalusian village", 5987261),
    "stone_mill":       ("Traditional Olive Oil Production", 5593361),
    "olive_press_workshop": ("Workshop with large stone wheel and old equipment for pressing olive oil", 2537351),
    "family_olive_grove": ("Video of drone of a family playfully runs through olive groves", 5682519),
    "cert_stamp":       ("Stamp with CERTIFIED text of flat style isolated on white background", 3982793),
    "dark_bottle_label": ("Hand holding a small dark green glass bottle with a red cap", 3717430),
    "tasting_reaction":  ("Woman Tastes Tea In A Glass And Smiles In Her House", 5079112),
    "hand_picking_olives": ("Close view of hand picking black olives from green tree branches", 6295653),
    "olive_harvest_people": ("People In Olive Orchard Collecting Fruit In The Harvest Season", 6127160),
    "bottle_pouring":   ("Pouring pure Extra Virgin Olive Oil", 606843),
    "olive_grove_hillside": ("Vast Olive Grove at Sunset", 6163120),
    # ---- expanded pool (150-200 unique target, round 2) ----
    "salad_pour1": ("Pouring Olive Oil on a Fresh Salad", 3522160),
    "salad_pour2": ("Tilting bottle pouring oil over tomato, onion, basil", 8663762),
    "salad_pour3": ("Pouring hand tilting olive oil bottle over salad on counter", 7402383),
    "cart_timelapse": ("An empty grocery cart drives a grocery supermarket, time lapse", 4244202),
    "cart_aisle": ("Shopping Cart in a Supermarket Aisle", 6154249),
    "aisle_colorful": ("A Vibrant Aisle of Colorful Products in a Modern Supermarket", 6451901),
    "handshake1": ("Hand shake, diversity partnership", 1233440),
    "handshake_boardroom": ("Business people handshake in boardroom, corporate partnership deal", 2466239),
    "handshake_negotiation": ("Executives Reach Handshake Across Meeting Table, Negotiation Scene", 7819113),
    "stock_decline1": ("Decreasing trend red line graph, stock market stagnation", 6590339),
    "stock_crash": ("Stock Market Crash Red Arrow Graph Going Down Into Recession", 4045992),
    "stock_down": ("Red candles stock market down, arrow drop", 6169589),
    "couple_cooking": ("Couple cooking in the kitchen", 2947891),
    "woman_pour_veg": ("Biracial woman pouring oil on vegetables in baking tray", 2717206),
    "senior_woman_salad": ("Happy senior woman mixing salad in sunny kitchen", 3169363),
    "man_olives": ("Relaxed man enjoys typical olives soaked in olive oil", 2508365),
    "spoon_olives": ("The olives are in a spoon, olive oil flows", 4312045),
    "lab_colorchange": ("A test tube with liquid changes color during a chemical reaction", 5272862),
    "vintage_factory": ("Archival footage, women working in factories, 1915", 98280),
    "harvest_process": ("Olive harvesting process", 6277209),
    "sunset_dolly": ("Dolly Shot Of Sun Setting Behind Olive Trees", 3322694),
    "spain_flag1": ("Waving Flag of Spain", 8567147),
    "spain_flag2": ("Spain national flag waving on flagpole", 3082160),
    "pruning_aerial": ("Aerial video of a man pruning olive tree with chainsaw", 1255745),
    "harvest_rack": ("Farmer harvesting olive with rack", 1848680),
    "harvest_electric_rake": ("Farmer harvesting olives using electric rake", 6052701),
    "spanish_harvest_hand": ("Spanish man harvesting olive trees in a sunny field by hand", 1060909),
    "spanish_harvest_dog": ("Typical Spanish man harvesting olives with his dog and old tools", 839038),
    "price_tags1": ("Price tags at a market stall with blurred fresh produce", 6143948),
    "price_tags2": ("Close-up of Price Tags at a Market Stall", 3520905),
    "price_tag3": ("Price tag for produce at a market", 5303140),
    "price_tag_grocery": ("Grocery store food price tag", 6008763),
    "branches_wind1": ("Branches with young green olives swinging in the wind", 222576),
    "branches_wind2": ("Spanish Olive Tree Branches Sway in Wind with Blue Sky", 742110),
    "branch_daylight": ("Close view of olive tree branch with leaves moving in daylight", 7504652),
    "counting_cash1": ("Employee hands counting earnings cash close up", 3015782),
    "counting_euros2": ("Hands counting new fifty euro bills, money stack", 989385),
    "counting_euros3": ("Woman counting fifty and one hundred euro bills", 6947932),
    "lab_testtubes1": ("A wooden rack holds test tubes filled with colorful liquids", 5712239),
    "lab_testtubes2": ("Scientist holding a plate of colorful test tubes with liquid solutions", 8182467),
    "lab_reagent": ("Chemist drips reagent into test tube with yellow liquid", 1747463),
    "forklift1": ("Forklift operator steering down warehouse aisle carrying pallet of boxes", 8378084),
    "forklift2": ("Warehouse Forklift Operation", 3862405),
    "boxes_shelves": ("Warehouse shelves filled with stacked cardboard boxes", 6684302),
    "forklift_boxes": ("A forklift carries stacked cardboard boxes along a warehouse aisle", 8410557),
    "reading_label1": ("African American Woman Reading Label on Olive Oil Bottle", 6935335),
    "reading_label2": ("Woman choosing products in supermarket, reading label on bottle of oil", 2362890),
    "golden_pour1": ("Close-up of Golden Olive Oil Pouring onto a Green Olive", 8006564),
    "golden_flow": ("Golden olive oil flowing smoothly during processing", 7291357),
    "bottle_pour2": ("Olive Oil Pouring from Bottle", 3878595),
    "cazorla_grove": ("Endless rows of olive trees around Cazorla village, Andalusia, Spain", 666966),
    "granada_groves": ("Trees and olive groves in green landscapes, Granada, Spain", 7055616),
    "restabal_village": ("Restábal, a white village in the Lecrín Valley, Granada, Spain", 6976685),
    "highway_olives": ("Aerial pan shot of a highway in the south of Spain surrounded by olive fields", 1106294),
    # ---- round 3 ----
    "checkout_scan1": ("Saleswoman scanning product at checkout counter in bright supermarket", 1755128),
    "checkout_mobile": ("Grocery Store Checkout with Mobile Payment", 5261444),
    "news_bankruptcy": ("Bankruptcy news headline in different articles", 2439631),
    "news_debtcrisis": ("Newspaper headlines saying Debt Crisis over stock market display", 1923745),
    "financial_crisis_text": ("Minimal cinematic financial crisis text animation", 8214134),
    "vintage_calculator1": ("Person Using an Old Red Calculator", 3750206),
    "adding_machine": ("Close up of printed numbers adding up with plus symbols on a paper roll", 7223545),
    "branch_bottles": ("Olive tree branches with small bottles of olive oil", 6643654),
    "oil_factory_worker1": ("Female worker working in oil factory", 1848714),
    "oil_factory_worker2": ("Female worker working in oil factory", 1848716),
    "oil_can_factory": ("Olive oil can being kept on table in oil factory", 2955969),
    "modern_oil_factory": ("Interior of modern natural oil factory", 6067479),
    "technicians_oil": ("Technicians examining olive oil", 1848743),
}

DOWNLOAD_URLS = {
    151138: "https://videocdn.cdnpk.net/videos/d7fe033e-4263-4364-b2db-a0c9639c6496/horizontal/downloads/4k.mp4?filename=997677_Woman_Shopper_3840x2160.mp4&user_id=102901249&token=exp=1785247408~id=102901249~hmac=17e7efdce4f2ecb8cad4fb32e709238f093b800f3163ba790a3af502aa48612c",
    687809: "https://videocdn.cdnpk.net/videos/c0e0ccc3-d109-4fc6-93c2-916264ab153f/horizontal/downloads/original.mp4?filename=1739877_Food_Onion_3840x2160.mp4&user_id=102901249&token=exp=1785247409~id=102901249~hmac=a62446820de749891a4be379f1e282c55e6f50a3136df0e63bf1324bf3989dc2",
    3472735: "https://videocdn.cdnpk.net/videos/9d2ffeb7-7287-5c9b-935b-d24177ce8d24/horizontal/downloads/original.mp4?filename=0_Olive_Oil_Bottles_3840x2160.mp4&user_id=102901249&token=exp=1785247411~id=102901249~hmac=7406859d4f4fb59c859afee5fd760d0de5b1833e1c097295447b90dd95c6f508",
    6264187: "https://videocdn.cdnpk.net/videos/cc5cebf9-db3d-5148-a5bd-c4cc26937c83/horizontal/downloads/original.mp4?filename=0_Olive_Oil_3840x2160.mp4&user_id=102901249&token=exp=1785247412~id=102901249~hmac=289e7ee4673fd321af220e73bdcccccdfb4e22baf6acb205201f197ac491fe26",
    8534240: "https://videocdn.cdnpk.net/videos/f03eb199-150d-55e5-a2c9-fd2d2163ba1f/horizontal/downloads/4k.mp4?filename=0_Magnifying_Glass_3840x2160.mp4&user_id=102901249&token=exp=1785247413~id=102901249~hmac=c913cfa5e79bbaf55210677d944d072ffc06424bd9a3c2bd34da0e52a0ddea67",
    1864061: "https://videocdn.cdnpk.net/videos/e5b149c4-77de-4f5c-b4ee-d1e7380d585c/horizontal/downloads/4k.mp4?filename=5351876_Coll_wavebreak_Laboratory_3840x2160.mp4&user_id=102901249&token=exp=1785247415~id=102901249~hmac=f2690007e7bdb8ce9ab6032a3fda4ad17292d1ba382a5fd88b7d14a3c0db9dc3",
    7917819: "https://videocdn.cdnpk.net/videos/2ebb5ed4-724f-551f-a446-03916f39f076/horizontal/downloads/original.mp4?filename=0_Grocery_Store_Supermarket_3840x2160.mp4&user_id=102901249&token=exp=1785247416~id=102901249~hmac=1ac3b0a2f6782ca723fa2856892855bfdae73f10ed972d148307af0d7845b30d",
    5177810: "https://videocdn.cdnpk.net/videos/437a48d7-993f-58ca-90db-f41b2d5e5dfb/horizontal/downloads/original.mp4?filename=0_Bottles_Oil_3840x2160.mp4&user_id=102901249&token=exp=1785247418~id=102901249~hmac=aecb4f1c98f363ca0f998a4b4ee806b6a2601ae0e8e5ccfbbe59dc266e55a8e8",
    7618572: "https://videocdn.cdnpk.net/videos/20a5e384-1ab4-5472-b787-ad3d4aa8bab3/horizontal/downloads/original.mp4?filename=0_Avatar_Network_4096x2160.mp4&user_id=102901249&token=exp=1785247419~id=102901249~hmac=04fdfcdeaecbaac57f8505d72498852845227f4ce1435d2081c65f74f7dba94a",
    3811676: "https://videocdn.cdnpk.net/videos/f49b1510-3928-5ba2-9d51-fecebd4195d1/horizontal/downloads/original.mp4?filename=0_Bottle_Plastic_Bottle_1920x1080.mp4&user_id=102901249&token=exp=1785247420~id=102901249~hmac=7d8ef336a750b2559c7150b739fa3c32161b8e4ec383257dc490ba48800e835a",
    6285: "https://videocdn.cdnpk.net/videos/965e9a60-5f19-4126-8f32-ec5a6607d1d2/horizontal/downloads/original.mp4?filename=7073_Euro_Money_2048x1080.mp4&user_id=102901249&token=exp=1785247422~id=102901249~hmac=965c72db0080681a320f7eb81d393bd437c9efab60891add5040b9ad703af358",
    5465997: "https://videocdn.cdnpk.net/videos/e49365cd-0e1f-5cb6-a350-74fc3203a4b3/horizontal/downloads/4k.mp4?filename=0_Olive_Oil_3840x2160.mp4&user_id=102901249&token=exp=1785247423~id=102901249~hmac=1102c480442c9a7ee3f2f6265d0ec7df992aecab3a37589c57305801a2e2fe35",
    5474555: "https://videocdn.cdnpk.net/videos/2b91cb8f-67a4-5776-a6f2-59d24d818447/horizontal/downloads/original.mp4?filename=0_Bottle_Opener_Bottle_Cap_1920x1080.mp4&user_id=102901249&token=exp=1785247424~id=102901249~hmac=13e90475b5d92df5901efd7cead115ffbd99e1a53903f64cd0d1607b54294667",
    5540725: "https://videocdn.cdnpk.net/videos/f61b3687-3477-5325-8550-2bdff962976b/horizontal/downloads/4k.mp4?filename=0_Bottling_Plant_Manufacturing_3840x2160.mp4&user_id=102901249&token=exp=1785248046~id=102901249~hmac=3d9ec94787effe8883e0746f93c6e88c622f47465f9100cbace5b776a9386799",
    5540686: "https://videocdn.cdnpk.net/videos/28d600e2-1c0d-5a86-aa34-912c2289bac2/horizontal/downloads/4k.mp4?filename=0_Bottling_Plant_Plastic_Bottles_3840x2160.mp4&user_id=102901249&token=exp=1785248047~id=102901249~hmac=4e7aeb94eccb90866f9f8272c39eded5339b9a5a846933ffe6798f43ed071e53",
    6956087: "https://videocdn.cdnpk.net/videos/61b1c7b7-3e52-5d6e-89a6-46fc9ce8e7ce/horizontal/downloads/original.mp4?filename=0_Olive_Grove_Olive_Trees_3840x2160.mp4&user_id=102901249&token=exp=1785248049~id=102901249~hmac=0e5ee34eab9ce34ef3e9a11d70d6ff09dc21fa342354a3fb67f653df6c5e36fc",
    6163120: "https://videocdn.cdnpk.net/videos/44a95d95-703d-5d75-a96c-d047f096b337/horizontal/downloads/4k.mp4?filename=0_Agriculture_Farm_3840x2160.mp4&user_id=102901249&token=exp=1785248050~id=102901249~hmac=92b102665598b9d1c84c4dabbe5545b5076b7dc4ecd9f95f650ef3d6eb4e9965",
    7249724: "https://videocdn.cdnpk.net/videos/fbfe242d-ca56-5009-a210-0379dabb6cac/horizontal/downloads/original.mp4?filename=0_Barrel_Industrial_3840x2160.mp4&user_id=102901249&token=exp=1785248052~id=102901249~hmac=e0a6e6dcc23dd854c66250ad0b88c81f59737f91de2cd3cdc91ff03116da6c83",
    2310359: "https://videocdn.cdnpk.net/videos/5f999d43-ab68-44c4-9a16-3abe9a26ccd5/horizontal/downloads/4k.mp4?filename=5812797_Coll_wavebreak_Bottles_3840x2160.mp4&user_id=102901249&token=exp=1785248053~id=102901249~hmac=29bba3e61a09872d6285a2c447a923151bbab45679171a76a386a21cc2ba87b8",
    5883923: "https://videocdn.cdnpk.net/videos/d6dab745-b699-52ea-a78b-a3387b9c2b60/horizontal/downloads/4k.mp4?filename=0_Container_Port_Cargo_Ships_3840x2020.mp4&user_id=102901249&token=exp=1785248055~id=102901249~hmac=bf2152201d4cd4703248f91bd9dfcaf4fb11945bb0ce2e539acb285bd0be473e",
    7558377: "https://videocdn.cdnpk.net/videos/b48af2b9-42ef-5f11-a2a2-b5f9460de25c/horizontal/downloads/original.mp4?filename=0_World_Map_Overlay_1920x1080.mp4&user_id=102901249&token=exp=1785248056~id=102901249~hmac=9e25699ddfb105ad033dc721f0057cf8292db2144e0bdaa81f71a76dc88b6241",
    6181813: "https://videocdn.cdnpk.net/videos/5bc6e637-d5df-5329-939b-70b5273de9fa/horizontal/downloads/1080p.mp4?filename=0_Grocery_Store_1920x1080.mp4&user_id=102901249&token=exp=1785248057~id=102901249~hmac=e5e968dd6fd3b48fc1c3c74f66dbc6654743afaee510f8f4c416f0288882f0f8",
    6977454: "https://videocdn.cdnpk.net/videos/dc27f940-b88d-5a39-bf51-3d1d4aa6764c/horizontal/downloads/original.mp4?filename=0_Olives_Jar_Of_Olives_3840x2160.mp4&user_id=102901249&token=exp=1785248059~id=102901249~hmac=6fd016a6aba7bf0346c8ce9680a2e920828010b0d767247369e30ff0fbf460ae",
    5987261: "https://videocdn.cdnpk.net/videos/c84569b6-13bb-5131-9d3d-c1b716d29585/horizontal/downloads/4k.mp4?filename=0_Andalusia_Spain_3840x2160.mp4&user_id=102901249&token=exp=1785248060~id=102901249~hmac=774d0d55b5f261e3536069503ca5063e3264b8fa5e9c9a639d4062cd2dcebe6c",
    5593361: "https://videocdn.cdnpk.net/videos/3b900971-4b7b-528c-971a-2f8cbec74351/horizontal/downloads/4k.mp4?filename=0_Olive_Oil_Production_Olive_Mill_3840x2160.mp4&user_id=102901249&token=exp=1785248061~id=102901249~hmac=96a5c2ebdc15179000972540416a64b7dd88633b5fddc95ab664ac586a3f397c",
    2537351: "https://videocdn.cdnpk.net/videos/28448975-da1f-45e3-946e-022e299225f3/horizontal/downloads/original.mp4?filename=5302999_Workshop_Large_Stone_Wheel_3840x2160.mp4&user_id=102901249&token=exp=1785248063~id=102901249~hmac=1bed493cdaeb530fd7f877e20024ce6d03aed8915f5fa3e1dea5c071ee517cdc",
    5682519: "https://videocdn.cdnpk.net/videos/a29e4f40-42ce-5d62-bb8a-70aa0285664b/horizontal/downloads/original.mp4?filename=0_Hovering_Away_3840x2160.mp4&user_id=102901249&token=exp=1785248064~id=102901249~hmac=6ce3969e76b58e8951c773fb6e5dcbe8eac14d98bd53234947976ec5378602dd",
    3982793: "https://delivery.gettyimages.com/downloads/1314159504?e=Nrkmd2RH96WrHHDKJ8QVWEL29kxDm6daIqd6hOt8klzomUrUwTLMQ14d5fiKEHRatW3fgFIcmH1ICjG_hVvz-mSqBu_Nrypu9wmc-s2KbnQjRUS8O_4--fGX4aWSxs4D&i=wGqlBaCtSn_TyU-VjIdgnw%3D%3D&k=49&c=IOMnITqnQtVCzIPUoLhnXQffpmye2GAn7G_n2FH6drI%3D",
    3717430: "https://videocdn.cdnpk.net/videos/10793104-cb33-543c-97f3-4fc9f8ec8696/horizontal/downloads/original.mp4?filename=0_Bottle_Glass_Bottle_3840x2160.mp4&user_id=102901249&token=exp=1785248067~id=102901249~hmac=f66c3c070c234b7b565a61ba8231e0a163694df2793d5232371ae45939409c00",
    5079112: "https://videocdn.cdnpk.net/videos/8bbdce2a-c4ef-5f75-ae41-d29a5c97c59e/horizontal/downloads/4k.mp4?filename=0_Woman_Tea_3840x2160.mp4&user_id=102901249&token=exp=1785248069~id=102901249~hmac=b35b3cad3c40cda57e981855e05c7cacd5364da911031f7032c48763919eb553",
    6295653: "https://videocdn.cdnpk.net/videos/5a836e84-1d81-5a02-9981-4d18d7e6a681/horizontal/downloads/4k.mp4?filename=0_Olive_Harvest_3840x2160.mp4&user_id=102901249&token=exp=1785248070~id=102901249~hmac=6939f0c32ec98cebf984f545d6e9014ccc5d4fb21dc1040f437594829e91e875",
    6127160: "https://videocdn.cdnpk.net/videos/bd10fc49-f3e0-5e18-a306-5a49d9bdcf4d/horizontal/downloads/4k.mp4?filename=0_People_Olive_3840x2160.mp4&user_id=102901249&token=exp=1785248071~id=102901249~hmac=f26da415599cbc8b63f53a4597e14b365a80ea3dd4432a5935f6017716c71fa7",
    606843: "https://videocdn.cdnpk.net/videos/d8e8475e-fe5c-4739-99f2-8b9bda0ffb7f/horizontal/downloads/4k.mp4?filename=1632451_Oil_Olive_3840x2160.mp4&user_id=102901249&token=exp=1785248073~id=102901249~hmac=7fb606bb049f6be9ee8f752b72a7cbfca3cae24de1b86dd516ed1477506d9383",
    **NEW_URLS,
}

# Per-section rotation of shots (cycled sentence by sentence for visual variety)
# Round 2: expanded pool (87 unique shots) redistributed to cut reuse from ~13x to ~5x avg.
SECTION_SHOTS = {
    "HOOK": ["buyer_choosing", "receipt_check", "shelf_oils", "premium_bottle", "magnifier_doc",
             "lab_pipette", "aisle_dolly", "bottles_row", "network_anim", "yellow_cap",
             "euros_rotating", "three_bottles", "opening_cap", "stock_decline1", "stock_crash",
             "stock_down", "cart_timelapse", "cart_aisle", "aisle_colorful", "counting_cash1",
             "news_debtcrisis", "financial_crisis_text", "checkout_scan1", "checkout_mobile"],
    "MARCA 7 - COOSUR": ["shelf_oils", "generic_shelf", "network_anim", "olive_grove_aerial",
                          "bottling_line", "warehouse_barrels", "conveyor_bottles", "olive_grove_sunset",
                          "yellow_cap", "magnifier_doc", "handshake_boardroom", "forklift1",
                          "boxes_shelves", "price_tags1", "aisle_colorful", "news_bankruptcy",
                          "modern_oil_factory"],
    "MARCA 6 - LA ESPANOLA": ["shelf_oils", "generic_shelf", "network_anim", "bottles_row",
                               "two_bottles_compare", "yellow_cap", "stone_mill", "magnifier_doc",
                               "olive_press_workshop", "vintage_factory", "reading_label2",
                               "price_tags2", "forklift2", "oil_factory_worker1", "vintage_calculator1"],
    "MARCA 5 - BORGES": ["shelf_oils", "port_containers", "world_map_port", "olive_grove_sunset",
                          "conveyor_bottles", "bottling_line", "lab_pipette", "premium_bottle",
                          "tasting_reaction", "bottle_pouring", "magnifier_doc", "handshake_negotiation",
                          "stock_crash", "forklift_boxes", "granada_groves"],
    "MARCA 4 - HACENDADO": ["shelf_oils", "generic_shelf", "network_anim", "bottling_line",
                             "yellow_cap", "magnifier_doc", "two_bottles_compare", "cart_aisle",
                             "cart_timelapse", "price_tag_grocery", "reading_label1", "aisle_colorful"],
    "MARCA 3 - YBARRA": ["andalusia_street", "yellow_cap", "two_bottles_compare", "stone_mill",
                          "lab_pipette", "magnifier_doc", "euros_rotating", "generic_shelf",
                          "spain_flag1", "branches_wind2", "cazorla_grove", "spanish_harvest_dog"],
    "MARCA 2 - HOJIBLANCA": ["shelf_oils", "network_anim", "bottling_line", "generic_shelf",
                              "yellow_cap", "premium_bottle", "world_map_port", "harvest_process",
                              "branches_wind1", "spanish_harvest_hand", "price_tag3", "oil_can_factory"],
    "MARCA 1 - CARBONELL": ["andalusia_street", "network_anim", "shelf_oils", "yellow_cap",
                             "two_bottles_compare", "stone_mill", "olive_press_workshop",
                             "bottle_pouring", "premium_bottle", "generic_shelf", "euros_rotating",
                             "vintage_factory", "stock_down", "handshake1", "sunset_dolly",
                             "counting_euros2", "adding_machine", "technicians_oil"],
    "MEJOR OPCION 3 - MELGAREJO": ["olive_grove_hillside", "hand_picking_olives", "cert_stamp",
                                     "olive_grove_aerial", "magnifier_doc", "pruning_aerial",
                                     "harvest_rack", "branch_daylight", "spoon_olives", "branch_bottles"],
    "MEJOR OPCION 2 - CASTILLO DE CANENA": ["olive_grove_sunset", "family_olive_grove", "cert_stamp",
                                              "dark_bottle_label", "magnifier_doc", "premium_bottle",
                                              "golden_pour1", "golden_flow", "lab_testtubes1", "man_olives"],
    "MEJOR OPCION 1 - NUNEZ DE PRADO": ["olive_harvest_people", "family_olive_grove", "bottle_pouring",
                                          "stone_mill", "cert_stamp", "olive_grove_hillside",
                                          "harvest_electric_rake", "restabal_village", "lab_testtubes2",
                                          "lab_reagent", "bottle_pour2", "oil_factory_worker2"],
    "CIERRE": ["three_bottles", "yellow_cap", "cert_stamp", "dark_bottle_label", "tasting_reaction",
               "bottle_pouring", "receipt_check", "spain_flag2", "counting_euros3", "salad_pour1",
               "salad_pour2", "salad_pour3", "couple_cooking", "woman_pour_veg", "senior_woman_salad",
               "highway_olives", "granada_groves"],
}

MOTIONS = [
    "ZOOM IN SLOW · 4s", "PAN RIGHT · 4s", "ZOOM OUT SLOW · 4s",
    "STATIC", "PAN LEFT · 4s", "DIAGONAL PAN + ZOOM IN · 4s",
    "ZOOM IN FAST · 2s", "PAN UP · 4s",
]


def split_sentences(text):
    text = text.replace("\n\n", " ").replace("\n", " ")
    text = re.sub(r"\s+", " ", text).strip()
    parts = re.split(r'(?<=[.!?])\s+(?=[A-ZÁÉÍÓÚÑ"¿¡])', text)
    return [p.strip() for p in parts if p.strip()]


def split_words(text, target=10):
    words = text.split()
    chunks, i = [], 0
    while i < len(words):
        chunk = words[i:i + target]
        if len(chunk) < 4 and chunks:
            chunks[-1] = chunks[-1] + " " + " ".join(chunk)
        else:
            chunks.append(" ".join(chunk))
        i += target
    return chunks


expanded = []  # (section, beat, text, shot_key, motion)
gcount = 0
motion_i = 0

for section, full_text in SECTIONS:
    shot_cycle = SECTION_SHOTS[section]
    sentences = split_sentences(full_text)
    for s_idx, sentence in enumerate(sentences):
        shot_key = shot_cycle[s_idx % len(shot_cycle)]
        for sub in split_words(sentence, target=10):
            gcount += 1
            motion = MOTIONS[motion_i % len(MOTIONS)]
            motion_i += 1
            expanded.append((section, gcount, sub, shot_key, motion))

TOTAL_WORDS = sum(len(t.split()) for _, t in SECTIONS)
NARRATION_SECONDS = 22 * 60 + 30  # 22:30 real ElevenLabs audio duration
SEC_PER_WORD = NARRATION_SECONDS / TOTAL_WORDS

timed = []
t = 0.0
for section, beat, text, shot_key, motion in expanded:
    words = len(text.split())
    dur = round(words * SEC_PER_WORD, 1)
    start = t
    end = t + dur
    t = end
    timed.append((section, beat, text, shot_key, motion, dur, start, end))


def fmt_tc(seconds):
    m = int(seconds // 60)
    s = seconds - m * 60
    return f"{m:02d}:{s:05.2f}"


print(f"Total real 3-5s beats: {gcount}")
print(f"Total words: {TOTAL_WORDS}")
print(f"Narracion real: {NARRATION_SECONDS}s -> {SEC_PER_WORD:.4f} s/palabra")
print(f"Duracion timeline calculada: {fmt_tc(t)} (objetivo 22:30.00)")

# ---------------- WRITE XLSX ----------------
wb = openpyxl.Workbook()

ws = wb.active
ws.title = "Beats"
headers = ["Sección", "Beat #", "Texto del guión", "Shot (descripción)", "Stock ID",
           "Duración (s)", "Inicio", "Fin", "CapCut Motion", "Ref. hoja Shots"]
ws.append(headers)
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="C62828")
    cell.alignment = Alignment(vertical="center")

for section, beat, text, shot_key, motion, dur, start, end in timed:
    title, sid = SHOTS[shot_key]
    ws.append([section, beat, text, title, sid, dur, fmt_tc(start), fmt_tc(end), motion, shot_key])

ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 8
ws.column_dimensions["C"].width = 50
ws.column_dimensions["D"].width = 45
ws.column_dimensions["E"].width = 10
ws.column_dimensions["F"].width = 12
ws.column_dimensions["G"].width = 10
ws.column_dimensions["H"].width = 10
ws.column_dimensions["I"].width = 20
ws.column_dimensions["J"].width = 20
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

ws2 = wb.create_sheet("Shots (unicos)")
ws2.append(["Ref.", "Titulo del clip", "Stock ID", "Licencia", "Enlace de descarga"])
for cell in ws2[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="C62828")

for key, (title, sid) in SHOTS.items():
    ws2.append([key, title, sid, "Premium (Freepik/Magnific)", DOWNLOAD_URLS.get(sid, "ENLACE CADUCADO - pedir de nuevo")])

ws2.column_dimensions["A"].width = 22
ws2.column_dimensions["B"].width = 70
ws2.column_dimensions["C"].width = 12
ws2.column_dimensions["D"].width = 12
ws2.column_dimensions["E"].width = 90

# ---------------- SHEET 3: Clip -> Beats (mapa inverso) ----------------
from collections import defaultdict
beats_by_shot = defaultdict(list)
for section, beat, text, shot_key, motion in expanded:
    beats_by_shot[shot_key].append(beat)

def compress_ranges(nums):
    nums = sorted(nums)
    ranges = []
    start = prev = nums[0]
    for n in nums[1:]:
        if n == prev + 1:
            prev = n
            continue
        ranges.append(f"{start}" if start == prev else f"{start}-{prev}")
        start = prev = n
    ranges.append(f"{start}" if start == prev else f"{start}-{prev}")
    return ", ".join(ranges)

ws3 = wb.create_sheet("Clip -> Beats")
ws3.append(["Ref.", "Titulo del clip", "Stock ID", "Nº de beats", "Beats asociados"])
for cell in ws3[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="C62828")

for key, (title, sid) in sorted(SHOTS.items(), key=lambda kv: -len(beats_by_shot.get(kv[0], []))):
    beats = beats_by_shot.get(key, [])
    if not beats:
        continue
    ws3.append([key, title, sid, len(beats), compress_ranges(beats)])

ws3.column_dimensions["A"].width = 22
ws3.column_dimensions["B"].width = 65
ws3.column_dimensions["C"].width = 12
ws3.column_dimensions["D"].width = 12
ws3.column_dimensions["E"].width = 60
for row in ws3.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

# ---------------- SHEET 4: Pendiente por buscar ----------------
ws4 = wb.create_sheet("Pendiente por buscar")
ws4.append(["Categoria", "Detalle"])
for cell in ws4[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="C62828")

pendientes = [
    ("Descargas bloqueadas por rate limit (Freepik)",
     "Olives in oil mills to make virgin olive oil - ID 977016"),
    ("Descargas bloqueadas por rate limit (Freepik)",
     "Pouring olives into machine at olive oil factory - ID 1348524"),
    ("Descargas bloqueadas por rate limit (Freepik)",
     "Low-angle EU flags outside European Commission - ID 5751633"),
    ("Descargas bloqueadas por rate limit (Freepik)",
     "European Commission headquarters EU flags slow motion - ID 5752936"),
    ("Descargas bloqueadas por rate limit (Freepik)",
     "Chef Drizzling Olive Oil and Balsamic Glaze Over Tomato - ID 8809408"),
    ("Clip con ID invalido (descartar)",
     "Magnifying glass 'The Fine Print' - ID 134130 devolvio 404, buscar alternativa"),
    ("Reutilizacion alta a revisar",
     "generic_shelf, shelf_oils, network_anim, magnifier_doc y yellow_cap superan 15-25 usos "
     "cada uno - son los primeros candidatos a sustituir si se sigue ampliando el pool"),
    ("Temas aun no buscados / cobertura floja",
     "Entrevista o testimonio de experto/productor a camara (aunque sea generico stock)"),
    ("Temas aun no buscados / cobertura floja",
     "Comparativa visual botella oscura vs botella clara explicando proteccion de la luz"),
    ("Temas aun no buscados / cobertura floja",
     "Tabla o grafico de categorias IOC (virgen extra / virgen / lampante) - infografia stock"),
    ("Temas aun no buscados / cobertura floja",
     "Fachada generica de supermercado tipo Mercadona/Carrefour (exterior de tienda, no interior)"),
    ("Temas aun no buscados / cobertura floja",
     "Laboratorio especifico de analisis sensorial de aceite (catadores con copas azules oficiales)"),
    ("Temas aun no buscados / cobertura floja",
     "Camion cisterna o camion de reparto de alimentacion en carretera (distribucion nacional)"),
    ("Temas aun no buscados / cobertura floja",
     "Cierre de fabrica / persiana bajada - para el tramo de crisis financiera de Deoleo/SOS Cuetara"),
]
for cat, detail in pendientes:
    ws4.append([cat, detail])

ws4.column_dimensions["A"].width = 32
ws4.column_dimensions["B"].width = 95
for row in ws4.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

wb.save("/tmp/claude-0/-home-user-Claudeeee/fed35260-2711-5769-ad05-7e136fd68906/scratchpad/Aceite_Oliva_PRODUCTION.xlsx")
print("Saved.")
print(f"Unique shots: {len(SHOTS)}")
