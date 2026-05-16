#!/usr/bin/env python3
"""Standalone VIDEO PROMPTS document for Video 2 (The Dopamine Trap).
Same layout style as the Video 1 video-prompts doc: split into parts,
BEAT n "segment" then Video Prompt."""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from build_dopamine_doc import BEATS  # single source of truth

TOTAL = len(BEATS)
CHUNK = 36
parts = []
i = 0
while i < TOTAL:
    parts.append((i + 1, min(i + CHUNK, TOTAL)))
    i += CHUNK

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(10.5)

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("CRAYON CAPITAL — CLONE SESSION · VIDEO 2")
r.bold = True
r.font.size = Pt(18)

s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run('STATE 9 — VIDEO PROMPTS')
r.bold = True
r.font.size = Pt(13)

s2 = doc.add_paragraph()
s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s2.add_run('"Por qué comprar se siente mejor que tener" · The Dopamine Trap')
r.italic = True
r.font.size = Pt(11)

intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
intro.add_run("Un video prompt para cada uno de los {} beats. "
              "Por volumen, en {} partes.".format(TOTAL, len(parts))).font.size = Pt(9.5)

for pidx, (a, b) in enumerate(parts, 1):
    doc.add_paragraph()
    ph = doc.add_paragraph()
    pr = ph.add_run("VIDEO PROMPTS — PARTE {} (Beats {}-{})".format(pidx, a, b))
    pr.bold = True
    pr.font.size = Pt(13)
    pr.font.color.rgb = RGBColor(0xB0, 0x2A, 0x2A)

    for n in range(a, b + 1):
        seg, scene, cam, light, mood, action, video = BEATS[n - 1]
        bp = doc.add_paragraph()
        bp.paragraph_format.space_before = Pt(8)
        br = bp.add_run("BEAT {}  ".format(n))
        br.bold = True
        br.font.size = Pt(11)
        br.font.color.rgb = RGBColor(0xB0, 0x2A, 0x2A)
        sr = bp.add_run('"{}"'.format(seg))
        sr.bold = True
        sr.italic = True
        sr.font.size = Pt(10.5)

        vp = doc.add_paragraph()
        lab = vp.add_run("Video Prompt: ")
        lab.bold = True
        vp.add_run(video)

    cp = doc.add_paragraph()
    cr = cp.add_run("PARTE {} COMPLETA — Beats {}-{} OK".format(pidx, a, b))
    cr.bold = True
    cr.font.color.rgb = RGBColor(0x1E, 0x7A, 0x33)

docx_path = "/home/user/Claudeeee/Dopamine_Trap_VIDEO_PROMPTS.docx"
doc.save(docx_path)
print("Saved DOCX:", docx_path, "|", TOTAL, "beats |", len(parts), "parts")
